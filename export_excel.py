"""Excel export for XPL wallet unlock data.

Generates a multi-sheet .xlsx workbook with:
  - Summary sheet (classification stats, totals)
  - Large Wallets sheet (detailed per-wallet breakdown)
  - All Wallets sheet (full distribution)
  - Data Sources sheet (where the data came from)
"""

from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

import config


# ── Color palette ──────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F2937")
SUBTITLE_FONT = Font(bold=True, size=11, color="4B5563")
NUM_FONT = Font(name="Consolas", size=10)
THIN_BORDER = Border(
    bottom=Side(style="thin", color="E5E7EB"),
)

BEHAVIOR_COLORS = {
    "heavy_seller": PatternFill("solid", fgColor="FEE2E2"),   # red-100
    "moderate_seller": PatternFill("solid", fgColor="FEF3C7"), # amber-100
    "light_seller": PatternFill("solid", fgColor="DBEAFE"),    # blue-100
    "holder": PatternFill("solid", fgColor="D1FAE5"),          # green-100
}


def export_excel(
    tracking_df: pd.DataFrame,
    stats: dict,
    all_wallets_df: pd.DataFrame = None,
    output_path: str = "xpl_unlock_report.xlsx",
    run_config: dict = None,
):
    """
    Write a styled Excel workbook with unlock tracking data.

    Parameters
    ----------
    tracking_df : DataFrame
        Large wallets with selling metrics.
    stats : dict
        Classification statistics from classify_wallets().
    all_wallets_df : DataFrame, optional
        Full wallet list (large + small).
    output_path : str
    run_config : dict, optional
        CLI / run configuration metadata.
    """
    wb = Workbook()

    _write_summary_sheet(wb, stats, tracking_df, run_config)
    _write_large_wallets_sheet(wb, tracking_df)
    if all_wallets_df is not None and not all_wallets_df.empty:
        _write_all_wallets_sheet(wb, all_wallets_df)
    _write_sources_sheet(wb, run_config)

    # Remove default empty sheet if it still exists
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(output_path)
    return output_path


# ── Sheet writers ──────────────────────────────────────────

def _write_summary_sheet(wb, stats, tracking_df, run_config):
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_properties.tabColor = "1F2937"

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = "XPL Wallet Unlock Tracker — Summary Report"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left")

    ws.merge_cells("A2:F2")
    ws["A2"] = f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A2"].font = Font(italic=True, color="9CA3AF", size=10)

    row = 4

    # Run config
    if run_config:
        ws.cell(row, 1, "RUN CONFIGURATION").font = SUBTITLE_FONT
        row += 1
        for key, val in run_config.items():
            ws.cell(row, 1, key.replace("_", " ").title()).font = Font(bold=True, size=10)
            ws.cell(row, 2, str(val)).font = Font(size=10)
            row += 1
        row += 1

    # Distribution stats
    ws.cell(row, 1, "DISTRIBUTION STATISTICS").font = SUBTITLE_FONT
    row += 1
    stat_items = [
        ("Total Wallets", stats.get("n_total", 0)),
        ("Large Wallets", f"{stats.get('n_large', 0)} ({stats.get('pct_large', 0)}%)"),
        ("Classification Method", stats.get("method", "")),
        ("Threshold", stats.get("threshold", "")),
        ("Cutoff Value (XPL)", f"{stats.get('cutoff_xpl', 0):,.2f}"),
        ("Mean Unlock", f"{stats.get('mean', 0):,.2f} XPL"),
        ("Median Unlock", f"{stats.get('median', 0):,.2f} XPL"),
        ("Std Deviation", f"{stats.get('std', 0):,.2f} XPL"),
        ("Min / Max", f"{stats.get('min', 0):,.2f} / {stats.get('max', 0):,.2f} XPL"),
        ("Q25 / Q75", f"{stats.get('q25', 0):,.2f} / {stats.get('q75', 0):,.2f} XPL"),
    ]
    for label, value in stat_items:
        ws.cell(row, 1, label).font = Font(bold=True, size=10)
        ws.cell(row, 2, str(value)).font = Font(size=10)
        row += 1

    row += 1

    # Behavior breakdown
    if not tracking_df.empty and "behavior" in tracking_df.columns:
        ws.cell(row, 1, "SELLING BEHAVIOR BREAKDOWN").font = SUBTITLE_FONT
        row += 1
        behavior_counts = tracking_df["behavior"].value_counts()
        for beh, count in behavior_counts.items():
            label = beh.replace("_", " ").title()
            c1 = ws.cell(row, 1, label)
            c1.font = Font(bold=True, size=10)
            c1.fill = BEHAVIOR_COLORS.get(beh, PatternFill())
            ws.cell(row, 2, count).font = Font(size=10)
            row += 1

        row += 1
        total_recv = tracking_df["amount_xpl"].sum()
        total_sold = tracking_df["amount_sold"].sum()
        total_rem = tracking_df["remaining"].sum()

        ws.cell(row, 1, "Total XPL Received").font = Font(bold=True, size=10)
        ws.cell(row, 2, f"{total_recv:,.2f}").font = NUM_FONT
        row += 1
        ws.cell(row, 1, "Total XPL Sold").font = Font(bold=True, size=10)
        pct_s = total_sold / total_recv * 100 if total_recv else 0
        ws.cell(row, 2, f"{total_sold:,.2f} ({pct_s:.1f}%)").font = NUM_FONT
        row += 1
        ws.cell(row, 1, "Total XPL Remaining").font = Font(bold=True, size=10)
        pct_r = total_rem / total_recv * 100 if total_recv else 0
        ws.cell(row, 2, f"{total_rem:,.2f} ({pct_r:.1f}%)").font = NUM_FONT

    # Column widths
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 35


def _write_large_wallets_sheet(wb, tracking_df):
    ws = wb.create_sheet("Large Wallets")
    ws.sheet_properties.tabColor = "EF4444"

    if tracking_df.empty:
        ws["A1"] = "No large wallets found."
        return

    headers = [
        ("Rank", 6),
        ("Address", 46),
        ("Received (XPL)", 20),
        ("Current Balance", 20),
        ("Amount Sold (XPL)", 20),
        ("% Sold", 10),
        ("Remaining (XPL)", 20),
        ("Behavior", 18),
        ("Data Source", 22),
    ]

    for col, (name, width) in enumerate(headers, 1):
        c = ws.cell(1, col, name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    for i, (_, row_data) in enumerate(tracking_df.iterrows()):
        r = i + 2
        ws.cell(r, 1, i + 1).alignment = Alignment(horizontal="center")
        ws.cell(r, 2, row_data["address"]).font = Font(name="Consolas", size=10)
        ws.cell(r, 3, round(row_data["amount_xpl"], 2)).number_format = '#,##0.00'
        ws.cell(r, 4, round(row_data.get("current_balance", 0), 2)).number_format = '#,##0.00'
        ws.cell(r, 5, round(row_data.get("amount_sold", 0), 2)).number_format = '#,##0.00'
        ws.cell(r, 6, round(row_data.get("pct_sold", 0), 1)).number_format = '0.0"%"'
        ws.cell(r, 7, round(row_data.get("remaining", 0), 2)).number_format = '#,##0.00'

        behavior = row_data.get("behavior", "")
        c_beh = ws.cell(r, 8, behavior.replace("_", " ").title())
        c_beh.fill = BEHAVIOR_COLORS.get(behavior, PatternFill())
        c_beh.alignment = Alignment(horizontal="center")

        source = row_data.get("source", "Plasma RPC / PlasmaScan")
        ws.cell(r, 9, source)

        # Zebra striping
        if i % 2 == 1:
            for col in range(1, len(headers) + 1):
                cell = ws.cell(r, col)
                if not cell.fill or cell.fill.fgColor is None or cell.fill.fgColor.rgb == "00000000":
                    cell.fill = PatternFill("solid", fgColor="F9FAFB")

    # Freeze top row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(tracking_df) + 1}"


def _write_all_wallets_sheet(wb, all_wallets_df):
    ws = wb.create_sheet("All Wallets")
    ws.sheet_properties.tabColor = "6366F1"

    headers = [
        ("Rank", 6),
        ("Address", 46),
        ("Amount (XPL)", 20),
        ("Classification", 16),
        ("Z-Score", 12),
    ]

    for col, (name, width) in enumerate(headers, 1):
        c = ws.cell(1, col, name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    sorted_df = all_wallets_df.sort_values("amount_xpl", ascending=False).reset_index(drop=True)
    for i, (_, row_data) in enumerate(sorted_df.iterrows()):
        r = i + 2
        ws.cell(r, 1, i + 1).alignment = Alignment(horizontal="center")
        ws.cell(r, 2, row_data["address"]).font = Font(name="Consolas", size=10)
        ws.cell(r, 3, round(row_data["amount_xpl"], 2)).number_format = '#,##0.00'
        ws.cell(r, 4, "Large" if row_data.get("is_large") else "Normal").alignment = Alignment(horizontal="center")
        z = row_data.get("z_score", 0)
        ws.cell(r, 5, round(z, 3) if z else 0).number_format = '0.000'

        if i % 2 == 1:
            for col in range(1, len(headers) + 1):
                ws.cell(r, col).fill = PatternFill("solid", fgColor="F9FAFB")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(sorted_df) + 1}"


def _write_sources_sheet(wb, run_config):
    ws = wb.create_sheet("Data Sources")
    ws.sheet_properties.tabColor = "10B981"

    ws.merge_cells("A1:C1")
    ws["A1"] = "Data Sources & Methodology"
    ws["A1"].font = TITLE_FONT

    row = 3
    sources = [
        ("Plasma RPC", config.RPC_URL,
         "Primary blockchain data — block scanning, balance queries, transaction history"),
        ("PlasmaScan API", config.PLASMASCAN_API_URL,
         "Etherscan-compatible API for indexed transaction lookups by address"),
        ("Arkham Intelligence", config.ARKHAM_API_URL,
         "Entity labeling, wallet identification, and on-chain analytics"),
        ("XPL Token Contract", config.XPL_CONTRACT,
         "XPL token on Plasma mainnet (ERC-20 compatible)"),
    ]

    headers = ["Source", "Endpoint / Address", "Description"]
    widths = [22, 50, 65]
    for col, (name, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row, col, name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = w
    row += 1

    for name, url, desc in sources:
        ws.cell(row, 1, name).font = Font(bold=True, size=10)
        ws.cell(row, 2, url).font = Font(name="Consolas", size=9, color="2563EB")
        ws.cell(row, 3, desc).font = Font(size=10)
        row += 1

    row += 2
    ws.cell(row, 1, "METHODOLOGY").font = SUBTITLE_FONT
    row += 1
    methodology = [
        "1. Wallet unlock amounts are collected by scanning Plasma blocks for large XPL transfers",
        "   from known vesting/treasury contracts during the configured date window.",
        "2. Wallets are classified as 'large' using a configurable statistical method (MAD, z-score,",
        "   IQR, percentile, or log-normal z-score) with a tunable threshold parameter.",
        "3. Current balances are queried from Plasma RPC to calculate selling behavior.",
        "4. Behavior categories: Heavy Seller (>=75% sold), Moderate Seller (40-75%),",
        "   Light Seller (10-40%), Holder (<10% sold).",
        f"5. Monthly ecosystem unlock: ~{config.MONTHLY_ECOSYSTEM_UNLOCK_XPL:,} XPL",
    ]
    for line in methodology:
        ws.cell(row, 1, line).font = Font(size=10)
        ws.merge_cells(f"A{row}:C{row}")
        row += 1
